import openpyxl
from openpyxl.utils import column_index_from_string


def parse_matrix_workbook(
    path: str,
    *,
    header_row: int,
    name_col: str,
    job_col: str,
    start_col: str,
    training_start_col: str,
    email_col: str | None = None,
) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    name_i = column_index_from_string(name_col)
    job_i = column_index_from_string(job_col)
    start_i = column_index_from_string(start_col)
    train_i = column_index_from_string(training_start_col)
    email_i = column_index_from_string(email_col) if email_col else None

    training_types: list[str] = []
    col = train_i
    while True:
        val = ws.cell(row=header_row, column=col).value
        if not val:
            break
        training_types.append(str(val).strip())
        col += 1

    people: list[dict] = []
    r = header_row + 1
    while True:
        name = ws.cell(row=r, column=name_i).value
        if not name:
            break
        start = ws.cell(row=r, column=start_i).value
        job = ws.cell(row=r, column=job_i).value
        row: dict = {
            "name": str(name).strip(),
            "job_title": str(job).strip() if job else None,
            "start_date": start.isoformat() if hasattr(start, "isoformat") else (str(start) if start else None),
        }
        if email_i:
            email_val = ws.cell(row=r, column=email_i).value
            row["email"] = str(email_val).strip().lower() if email_val else None
        people.append(row)
        r += 1

    return {"people": people, "training_types": training_types}
