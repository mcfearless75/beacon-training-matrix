import openpyxl

from beacon.importer import parse_matrix_workbook


def make_sample_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Job Title"
    ws["C1"] = "Start Date"
    ws["D1"] = "Company Induction"
    ws["E1"] = "First Aid"
    ws["A2"] = "Alice Smith"
    ws["B2"] = "Operative"
    ws["C2"] = "2024-01-10"
    ws["A3"] = "Bob Jones"
    ws["B3"] = "Supervisor"
    ws["C3"] = "2023-05-01"
    wb.save(path)


def test_parse_matrix_workbook(tmp_path):
    p = tmp_path / "m.xlsx"
    make_sample_workbook(p)
    result = parse_matrix_workbook(
        str(p),
        header_row=1,
        name_col="A",
        job_col="B",
        start_col="C",
        training_start_col="D",
    )
    assert {"Alice Smith", "Bob Jones"} == {p["name"] for p in result["people"]}
    assert {"Company Induction", "First Aid"} == set(result["training_types"])


def test_parse_matrix_workbook_handles_empty(tmp_path):
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Job Title"
    ws["C1"] = "Start Date"
    wb.save(p)
    result = parse_matrix_workbook(
        str(p),
        header_row=1,
        name_col="A",
        job_col="B",
        start_col="C",
        training_start_col="D",
    )
    assert result["people"] == []
    assert result["training_types"] == []
